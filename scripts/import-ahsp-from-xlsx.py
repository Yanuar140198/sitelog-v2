"""
Import AHSP data from D:/FORM/AHSP_NEW_4.xlsx into normalized JSON.

Output: D:/sitelog-v2/scripts/ahsp-seed.json

Run:
    D:/FORM/.venv/Scripts/python.exe D:/sitelog-v2/scripts/import-ahsp-from-xlsx.py

The Bina Marga AHSP workbook has 3 reference sheets:
  - PARAMETER : HSD catalog (resource master) split into A=tenaga, B=bahan, C=peralatan.
  - INPUT     : global default assumptions (not used here — values already baked per block).
  - CALCULATION : 58 AHSP blocks. Each block:
      * Header row with `Satuan: <unit>` in col H, identifier in col B
      * Section I  "I. INPUT & ASUMSI"           — input variables
      * Section II "II. PRODUKTIVITAS & KOEFISIEN" — derived koefisien
      * Section III "III. ANALISA HARGA SATUAN"    — TENAGA / BAHAN / PERALATAN line items
"""
from __future__ import annotations

import io
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

XLSX_PATH = Path('D:/FORM/AHSP_NEW_4.xlsx')
OUT_PATH = Path('D:/sitelog-v2/scripts/ahsp-seed.json')

CATEGORY_MAP = {
    'A. HSD TENAGA KERJA': 'tenaga',
    'B. HSD BAHAN / MATERIAL': 'bahan',
    'C. HSD PERALATAN': 'peralatan',
}

KODE_REGEX = re.compile(r'\[([A-Za-z0-9-]+)\]')
SATUAN_REGEX = re.compile(r'Satuan:\s*(.+)$', re.IGNORECASE)


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s or s == '-':
        return None
    try:
        return float(s)
    except ValueError:
        return None


def text(v):
    if v is None:
        return ''
    return str(v).strip()


def parse_parameter(ws) -> list[dict]:
    """Walk PARAMETER sheet, picking up section headers and resource rows.

    Note: codes may collide across sections (e.g. L01 appears in both A and C);
    we keep both — caller can disambiguate by (kode, category).
    """
    resources: list[dict] = []
    current_cat: str | None = None
    in_table = False

    for row in ws.iter_rows(values_only=True):
        cells = list(row[:8]) + [None] * (8 - min(8, len(row)))
        joined = ' '.join(text(c) for c in cells if c is not None)

        # Detect section header (e.g. "A. HSD TENAGA KERJA")
        for header, cat in CATEGORY_MAP.items():
            if header in joined:
                current_cat = cat
                in_table = False
                break
        else:
            # Detect table header row
            if current_cat and text(cells[1]) == 'No.' and text(cells[2]) == 'Kode':
                in_table = True
                continue

            if in_table and current_cat:
                no, kode, uraian, satuan, hsd, catatan = cells[1], cells[2], cells[3], cells[4], cells[5], cells[6]
                if text(kode) and text(uraian) and num(hsd) is not None:
                    resources.append({
                        'kode': text(kode),
                        'nama': text(uraian),
                        'category': current_cat,
                        'satuan': text(satuan) or '-',
                        'hsd': num(hsd) or 0.0,
                        'catatan': text(catatan),
                    })
                elif not text(kode) and not text(uraian):
                    # blank row terminates the sub-table; wait for next header
                    in_table = False

    return resources


def col_b_to_kode(b: str) -> str:
    """Normalize col B header value into a unique DB kode.

    Examples:
        'AHSP #1'      -> 'AHSP-1'
        'CL-LC-001'    -> 'CL-LC-001'
        'EW-TOP-004'   -> 'EW-TOP-004'
    """
    s = b.strip()
    m = re.match(r'^AHSP\s*#\s*(\d+)', s)
    if m:
        return f'AHSP-{m.group(1)}'
    return s


def derive_label(b: str, c: str) -> tuple[str, str]:
    """Return (label, jenis) from col B and col C of the header row.

    For 'AHSP #1' / '[EI-311] Item: 3.1 (1) ...' -> ('EI-311', '3.1 (1) ...')
    For 'CL-LC-001' / 'Land Clearing' -> ('CL-LC-001', 'Land Clearing')
    """
    b = b.strip()
    c = c.strip()
    if b.startswith('AHSP'):
        m = KODE_REGEX.search(c)
        if m:
            label = m.group(1)
            # everything after the bracket — strip 'Item: ' prefix
            after = c[m.end():].strip()
            after = re.sub(r'^Item:\s*', '', after).strip()
            return (label or b, after or b)
        return (b, c or b)
    return (b, c or b)


def parse_calculation(ws) -> list[dict]:
    """Parse CALCULATION sheet into a list of AHSP block dicts."""
    rows = list(ws.iter_rows(values_only=True))

    # Find all block starts: rows where col H starts with 'Satuan:'
    block_starts: list[int] = []
    for i, row in enumerate(rows):
        cells = list(row[:9]) + [None] * (9 - min(9, len(row)))
        h = text(cells[7])
        if h.startswith('Satuan:'):
            block_starts.append(i)

    items: list[dict] = []
    seen_kodes: dict[str, int] = {}

    for idx, start in enumerate(block_starts):
        end = block_starts[idx + 1] if idx + 1 < len(block_starts) else len(rows)
        header = list(rows[start][:9]) + [None] * (9 - min(9, len(rows[start])))
        b = text(header[1])
        c = text(header[2])
        h = text(header[7])

        satuan_match = SATUAN_REGEX.search(h)
        satuan = satuan_match.group(1).strip() if satuan_match else ''

        kode = col_b_to_kode(b)
        # disambiguate accidental dupes
        if kode in seen_kodes:
            seen_kodes[kode] += 1
            kode = f'{kode}-v{seen_kodes[kode]}'
        else:
            seen_kodes[kode] = 1
        label, jenis = derive_label(b, c)

        item = {
            'kode': kode,
            'sourceKode': b,
            'label': label,
            'jenis': jenis,
            'satuan': satuan,
            'ohpPct': 0.0,
            'inputs': [],
            'koefisien': [],
            'resources': [],
            'abcSubtotal': 0.0,
            'ohpAmount': 0.0,
            'unitRate': 0.0,
        }

        # Walk rows from start+1 to end, classifying by section
        section = None  # 'input' | 'koef' | 'analisa'
        analisa_cat = None  # 'tenaga' | 'bahan' | 'peralatan'
        in_resource_table = False
        ordinal_in = 0
        ordinal_k = 0
        ordinal_r = {'tenaga': 0, 'bahan': 0, 'peralatan': 0}

        for r in range(start + 1, end):
            row = list(rows[r][:9]) + [None] * (9 - min(9, len(rows[r])))
            joined = ' '.join(text(x) for x in row if x is not None)

            if 'I. INPUT' in joined:
                section = 'input'
                in_resource_table = False
                continue
            if 'II. PRODUKTIVITAS' in joined or 'II. PRODUKTIVITAS & KOEFISIEN' in joined:
                section = 'koef'
                in_resource_table = False
                continue
            if 'III. ANALISA' in joined:
                section = 'analisa'
                in_resource_table = False
                continue

            # In section III, watch for category sub-headers and totals.
            # Column layout (1-indexed letters / 0-indexed values):
            #   B(1)=No.  C(2)=Kode  D(3)=Komponen  E(4)=Koefisien  F(5)=Satuan
            #   G(6)=HSD  H(7)=Total  I(8)=Keterangan
            # Subtotal/summary rows put the label in D(3) and the value in H(7).
            if section == 'analisa':
                a = text(row[1])  # 'A.' / 'B.' / 'C.' / 'D.' / 'E.' / 'F.'
                ccol = text(row[2])
                dcol = text(row[3])
                if a == 'A.' and (ccol.startswith('TENAGA') or dcol.startswith('TENAGA')):
                    analisa_cat = 'tenaga'
                    in_resource_table = True
                    continue
                if a == 'B.' and (ccol.startswith('BAHAN') or dcol.startswith('BAHAN')):
                    analisa_cat = 'bahan'
                    in_resource_table = True
                    continue
                if a == 'C.' and (ccol.startswith('PERALATAN') or dcol.startswith('PERALATAN')):
                    analisa_cat = 'peralatan'
                    in_resource_table = True
                    continue
                if a == 'D.':
                    item['abcSubtotal'] = num(row[7]) or 0.0
                    in_resource_table = False
                    continue
                if a == 'E.':
                    pct = num(row[4])  # OHP % is in E (col index 4)
                    item['ohpPct'] = pct or 0.0
                    item['ohpAmount'] = num(row[7]) or 0.0
                    in_resource_table = False
                    continue
                if a == 'F.':
                    item['unitRate'] = num(row[7]) or 0.0
                    in_resource_table = False
                    continue

                # Skip header row (No. | Kode | Komponen ...)
                if a == 'No.' and ccol == 'Kode':
                    continue

                # "JUMLAH HARGA ... (X)" subtotal lines (label lives in col D)
                if dcol.startswith('JUMLAH HARGA'):
                    continue
                # "Tidak ada bahan" placeholder
                if a == '-' or ccol.startswith('Tidak ada') or dcol.startswith('Tidak ada'):
                    continue

                # Actual resource line: cells: [No | Kode | Komponen | Koef | Satuan | HSD | Total | Keterangan]
                if in_resource_table and analisa_cat:
                    resource_code = ccol
                    uraian = dcol
                    koef = num(row[4])
                    sat = text(row[5])
                    hsd = num(row[6])
                    total = num(row[7])
                    if resource_code and uraian and koef is not None and hsd is not None:
                        ordinal_r[analisa_cat] += 1
                        item['resources'].append({
                            'category': analisa_cat,
                            'ordinal': ordinal_r[analisa_cat],
                            'kode': resource_code,
                            'uraian': uraian,
                            'koefisien': koef,
                            'satuan': sat,
                            'hsd': hsd,
                            'subtotal': total if total is not None else koef * hsd,
                        })
                continue

            # Sections I and II share the same column layout:
            # [Kode | Variable | Uraian | Nilai | Satuan | Sumber/Formula]
            kcol = text(row[1])
            vcol = text(row[2])
            ucol = text(row[3])
            ncol = num(row[4])
            scol = text(row[5])
            fcol = text(row[6])

            # Skip table-header rows
            if kcol == 'Kode' and vcol == 'Variable':
                continue
            # Skip warning rows
            if '⚠' in joined or 'CHECK REQUIRED' in joined:
                continue

            if not kcol and not ucol:
                continue

            if section == 'input':
                ordinal_in += 1
                item['inputs'].append({
                    'ordinal': ordinal_in,
                    'kode': kcol,
                    'variable': vcol if vcol and vcol != '-' else None,
                    'uraian': ucol,
                    'nilai': ncol,
                    'satuan': scol,
                    'sumber': fcol,
                })
            elif section == 'koef':
                ordinal_k += 1
                item['koefisien'].append({
                    'ordinal': ordinal_k,
                    'kode': kcol,
                    'variable': vcol if vcol and vcol != '-' else None,
                    'uraian': ucol,
                    'nilai': ncol,
                    'satuan': scol,
                    'formula': fcol,
                })

        items.append(item)

    return items


def main() -> None:
    print(f'Loading {XLSX_PATH} ...')
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True, read_only=True)

    resources = parse_parameter(wb['PARAMETER'])
    items = parse_calculation(wb['CALCULATION'])

    seed = {'resources': resources, 'items': items}

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(json.dumps(seed, indent=2, ensure_ascii=False), encoding='utf-8')

    total_resource_lines = sum(len(i['resources']) for i in items)
    total_inputs = sum(len(i['inputs']) for i in items)
    total_koefs = sum(len(i['koefisien']) for i in items)

    print('--- Import summary ---')
    print(f'Resources in catalog : {len(resources)}')
    print(f'AHSP items           : {len(items)}')
    print(f'Total input rows     : {total_inputs}')
    print(f'Total koefisien rows : {total_koefs}')
    print(f'Total resource lines : {total_resource_lines}')
    print(f'Output               : {OUT_PATH}')

    # Sanity check: pick EI-311 / AHSP-1
    sample = next((i for i in items if i['kode'] == 'AHSP-1'), None)
    if sample:
        print(f"\nSample AHSP-1 ({sample['label']}): rate={sample['unitRate']:.4f}  expected≈38102.44")
        for r in sample['resources']:
            print(f"  {r['category']:9s} {r['kode']:6s} {r['uraian'][:30]:30s} koef={r['koefisien']:.6f} hsd={r['hsd']:.2f} = {r['subtotal']:.2f}")


if __name__ == '__main__':
    main()
